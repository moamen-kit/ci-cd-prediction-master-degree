"""Rebuild the thesis .docx from the corrected markdown.

Uses the existing document as a style template: its styles.xml, section setup
and page geometry are preserved, the body is cleared, and the corrected content
is written back using the same named styles.
"""
import re, sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

MD   = Path(sys.argv[1])
TPL  = Path(sys.argv[2])
OUT  = Path(sys.argv[3])
FIGS = Path("cicd-failure-prediction/figures")

doc = docx.Document(str(TPL))
body = doc.element.body
for child in list(body):                      # clear body, keep sectPr
    if not child.tag.endswith('}sectPr'):
        body.remove(child)

styles = {s.name for s in doc.styles}

def _usable(name):
    """A style can be listed yet unresolvable (latent styles). Prove it works."""
    if name not in styles:
        return False
    try:
        doc.styles[name]
        return True
    except KeyError:
        return False

_ok = {}
def style_or(name, fallback="Normal"):
    if name not in _ok:
        _ok[name] = _usable(name)
    return name if _ok[name] else fallback

def add_heading_para(level, text):
    """Heading with the named style, or a manually formatted fallback."""
    st = style_or(f"Heading {level}")
    if st != "Normal":
        p = doc.add_paragraph(style=st)
        add_runs(p, text)
        return p
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); r.bold = True
    r.font.size = Pt({1: 18, 2: 15, 3: 13, 4: 11}.get(level, 11))
    return p

INLINE = re.compile(r'(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)', re.S)

def add_runs(par, text):
    """Render **bold**, *italic* and `code` into runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(9.5)
        else:
            par.add_run(part)

lines = MD.read_text(encoding="utf-8").split("\n")
i = 0
stats = dict(h=0, p=0, t=0, img=0, cap=0, li=0)

while i < len(lines):
    ln = lines[i]; s = ln.strip()

    if not s or s == "---":
        i += 1; continue

    # ---- image ------------------------------------------------------
    m = re.match(r'!\[([^\]]*)\]\((figures/[^)]+)\)', s)
    if m:
        path = FIGS / Path(m.group(2)).name
        if path.exists():
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Inches(6.0))
            stats["img"] += 1
        i += 1; continue

    # ---- heading ----------------------------------------------------
    m = re.match(r'^(#{1,6})\s+(.*)$', s)
    if m:
        lvl = min(len(m.group(1)), 4)
        add_heading_para(lvl, m.group(2)); stats["h"] += 1
        i += 1; continue

    # ---- table ------------------------------------------------------
    if s.startswith("|") and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i+1].strip()):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            r = lines[i].strip()
            if not re.match(r'^\|[\s:|-]+\|$', r):
                rows.append([c.strip() for c in r.strip("|").split("|")])
            i += 1
        if rows:
            ncols = max(len(r) for r in rows)
            tbl = doc.add_table(rows=0, cols=ncols)
            tbl.style = style_or("Table Grid", "Table")
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                cells = tbl.add_row().cells
                for ci in range(ncols):
                    txt = row[ci] if ci < len(row) else ""
                    cell = cells[ci]
                    cell.paragraphs[0].text = ""
                    add_runs(cell.paragraphs[0], txt)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
                        if ri == 0: run.bold = True
            stats["t"] += 1
        continue

    # ---- figure caption (whole line italic) --------------------------
    if s.startswith("*") and s.endswith("*") and not s.startswith("**") and len(s) > 40:
        p = doc.add_paragraph(style=style_or("Caption"))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, s); stats["cap"] += 1
        i += 1; continue

    # ---- blockquote --------------------------------------------------
    if s.startswith(">"):
        buf = []
        while i < len(lines) and lines[i].strip().startswith(">"):
            buf.append(lines[i].strip().lstrip(">").strip()); i += 1
        p = doc.add_paragraph(style=style_or("Quote"))
        p.paragraph_format.left_indent = Inches(0.4)
        add_runs(p, " ".join(x for x in buf if x))
        continue

    # ---- list --------------------------------------------------------
    m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', ln)
    if m:
        ordered = not m.group(2) in ("-", "*")
        st = style_or("List Number" if ordered else "List Bullet", "Normal")
        p = doc.add_paragraph(style=st)
        if st == "Normal":
            p.paragraph_format.left_indent = Inches(0.35)
        add_runs(p, m.group(3)); stats["li"] += 1
        i += 1; continue

    # ---- paragraph ----------------------------------------------------
    buf = []
    while i < len(lines):
        cur = lines[i].strip()
        if (not cur or cur == "---" or cur.startswith("|") or cur.startswith("#")
                or cur.startswith(">") or cur.startswith("![")
                or re.match(r'^(\s*)([-*]|\d+\.)\s+', lines[i])):
            break
        # A line that is entirely bold is a standalone line (title blocks),
        # not a soft-wrapped continuation of the paragraph.
        standalone = bool(re.fullmatch(r'\*\*.+\*\*', cur))
        if standalone and buf:
            break
        buf.append(cur); i += 1
        if standalone:
            break
    if buf:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        add_runs(p, " ".join(buf)); stats["p"] += 1

doc.save(str(OUT))
print("headings %(h)d · paragraphs %(p)d · tables %(t)d · images %(img)d · captions %(cap)d · list items %(li)d" % stats)
